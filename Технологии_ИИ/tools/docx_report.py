"""Сборка отчётов .docx в оформлении отчётов по лабам 1-2.

Times New Roman 14, полуторный интервал, абзацный отступ 1,25 см,
«Таблица N - ...» над таблицей, «Рисунок N - ...» под рисунком,
номер страницы внизу по центру.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = 'Times New Roman'
MONO = 'Courier New'


def num(x, nd=2):
    """Число в русской записи: пробел - разделитель тысяч, запятая - дробной части."""
    if isinstance(x, (int,)) or (isinstance(x, float) and nd == 0):
        s = f'{x:,.0f}'
    else:
        s = f'{x:,.{nd}f}'
    return s.replace(',', ' ').replace('.', ',')


def _set_font(run, name=FONT, size=14, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _page_number(paragraph):
    run = paragraph.add_run()
    _set_font(run, size=12)
    for tag, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        run._r.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        self.t = 0  # счётчик таблиц
        self.f = 0  # счётчик рисунков
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = self.doc.styles['Normal']
        st.font.name = FONT
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.5
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _page_number(fp)

    # --- текст -------------------------------------------------------------
    def _para(self, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, spacing=1.5,
              before=0, after=0, keep=False):
        p = self.doc.add_paragraph()
        f = p.paragraph_format
        p.alignment = align
        f.first_line_indent = Cm(indent) if indent else None
        f.line_spacing = spacing
        f.space_before, f.space_after = Pt(before), Pt(after)
        f.keep_with_next = keep
        f.widow_control = True
        return p

    def _runs(self, p, text, size=14, bold=False):
        """Текст с поддержкой **полужирного** и `моноширинного` фрагментов."""
        import re
        for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if not part:
                continue
            if part.startswith('**'):
                _set_font(p.add_run(part[2:-2]), size=size, bold=True)
            elif part.startswith('`'):
                _set_font(p.add_run(part[1:-1]), name=MONO, size=size - 2, bold=bold)
            else:
                _set_font(p.add_run(part), size=size, bold=bold)
        return p

    def section(self, title, new_page=False):
        if new_page and len(self.doc.paragraphs) > 0:
            self.page_break()
        p = self._para(align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, before=12, after=6, keep=True)
        _set_font(p.add_run(title), bold=True)

    def subsection(self, title):
        p = self._para(align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, before=8, after=4, keep=True)
        _set_font(p.add_run(title), bold=True)

    def p(self, text, keep=False):
        return self._runs(self._para(keep=keep), text)

    def items(self, items, numbered=False, dash='-'):
        for i, it in enumerate(items, 1):
            p = self._para(indent=None)
            f = p.paragraph_format
            f.left_indent = Cm(1.25 + 0.75)
            f.first_line_indent = Cm(-0.75)
            mark = f'{i}. ' if numbered else f'{dash} '
            self._runs(p, mark + it)

    def code(self, text, size=9):
        lines = text.strip('\n').split('\n')
        for i, line in enumerate(lines):
            p = self._para(align=WD_ALIGN_PARAGRAPH.LEFT, indent=None, spacing=1.0,
                           before=6 if i == 0 else 0,
                           after=6 if i == len(lines) - 1 else 0,
                           keep=i == 0 or (len(lines) <= 8 and i < len(lines) - 1))
            p.paragraph_format.left_indent = Cm(0.5)
            _set_font(p.add_run(line if line else ' '), name=MONO, size=size)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # --- таблицы -----------------------------------------------------------
    def table(self, caption, header, rows, widths=None, size=12, align=None):
        """align - строка из 'l','c','r' для столбцов."""
        self.t += 1
        p = self._para(align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, spacing=1.0,
                       before=6, after=3, keep=True)
        _set_font(p.add_run(f'Таблица {self.t} - {caption}'), size=size)
        tb = self.doc.add_table(rows=1, cols=len(header))
        tb.style = 'Table Grid'
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        amap = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER,
                'r': WD_ALIGN_PARAGRAPH.RIGHT}

        def fill(cell, text, bold=False, a='l'):
            cell.text = ''
            par = cell.paragraphs[0]
            par.alignment = amap[a]
            pf = par.paragraph_format
            pf.line_spacing, pf.first_line_indent = 1.0, None
            pf.space_before = pf.space_after = Pt(1)
            self._runs(par, str(text), size=size, bold=bold)

        for j, h in enumerate(header):
            fill(tb.rows[0].cells[j], h, bold=True, a='c')
        # шапка повторяется на каждой странице
        trpr = tb.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        th.set(qn('w:val'), 'true')
        trpr.append(th)
        for r in rows:
            row = tb.add_row()
            cs = OxmlElement('w:cantSplit')
            cs.set(qn('w:val'), 'true')
            row._tr.get_or_add_trPr().append(cs)
            cells = row.cells
            for j, v in enumerate(r):
                fill(cells[j], v, a=(align[j] if align else 'l'))
        if widths:
            tb.autofit = False
            grid = tb._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn('w:gridCol'))):
                gc.set(qn('w:w'), str(int(Cm(widths[j]).twips)))
            for row in tb.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        # небольшой отступ после таблицы
        self._para(spacing=1.0)
        return tb

    # --- рисунки -----------------------------------------------------------
    def figure(self, path, caption, width=16.0):
        self.f += 1
        p = self._para(align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, spacing=1.0,
                       before=6, keep=True)
        p.add_run().add_picture(str(path), width=Cm(width))
        c = self._para(align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, spacing=1.0,
                       before=3, after=8)
        _set_font(c.add_run(f'Рисунок {self.f} - {caption}'))
        return self.f

    # --- сохранение --------------------------------------------------------
    def save(self, path, pdf=True):
        path = Path(path)
        self.doc.save(path)
        if pdf:
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf',
                            '--outdir', str(path.parent), str(path)],
                           check=True, capture_output=True)
        return path
